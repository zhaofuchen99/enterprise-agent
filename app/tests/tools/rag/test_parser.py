"""文档解析（详细设计 11.2 / 11.3 的清洗规则）。

四种格式各测各的，因为它们能提供的信息量本来就不同（见 `parser.py` 模块 docstring）。

**PDF 用真文件测**：解析 PDF 的主要难点全在 `pdfplumber` 的实际输出形态上
（页眉页脚是画上去的、表格是独立对象、跨页要靠 `repeatRows`），
拿一个假对象去测只能证明我们自己的代码自洽。所以这里用 `reportlab`
现造一份结构真实的 PDF——它与 `scripts/gen_corpus.py` 用的是同一条排版路径。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import docx
import pytest
from docx.shared import Pt

from app.core.errors import AgentError, ErrorCode
from app.tools.rag.parser import (
    Block,
    ParsedDocument,
    TableBlock,
    _detect_page_furniture,
    _inherit_captions,
    _repeated,
    detect_format,
    parse_document,
)


def _write(path: Path, content: str) -> Path:
    path.write_text(content, encoding="utf-8")
    return path


# ------------------------------------------------------------------ 格式判定


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("a.pdf", "pdf"),
        ("a.PDF", "pdf"),
        ("a.docx", "docx"),
        ("a.md", "md"),
        ("a.markdown", "md"),
        ("a.txt", "txt"),
    ],
)
def test_format_detection(name: str, expected: str) -> None:
    assert detect_format(Path(name)) == expected


def test_unsupported_suffix_is_rejected() -> None:
    """不认得的扩展名**报错而不是猜**。

    猜的后果是拿一个 PDF 解析器去读 xlsx，然后报一个和真实原因无关的错。
    """
    with pytest.raises(AgentError) as excinfo:
        detect_format(Path("薪资表.xlsx"))

    assert excinfo.value.code == ErrorCode.INVALID_ARGUMENT
    assert "xlsx" in excinfo.value.message


def test_missing_file_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(AgentError) as excinfo:
        parse_document(tmp_path / "不存在.md")

    assert excinfo.value.code == ErrorCode.INVALID_ARGUMENT


# ------------------------------------------------------------------ Markdown


_MARKDOWN = """\
## 净销售额指标口径说明

**指标信息**
| 项目 | 内容 |
|---|---|
| 指标编码 | net_sales |
| 版本 | v1.2 |

## 一、指标定义

净销售额指**扣除让利与退回后**的收入金额。

## 二、计算公式

净销售额 = 含税销售额 − 折扣金额 − 退货金额

---

## 六、版本变更记录

**版本历史**
| 版本 | 日期 |
|---|---|
| v1.0 | 2024-01-01 |
"""


def test_markdown_reads_heading_levels(tmp_path: Path) -> None:
    parsed = parse_document(_write(tmp_path / "m.md", _MARKDOWN))

    headings = [(b.level, b.text) for b in parsed.blocks if b.kind == "HEADING"]

    # `六、版本变更记录` 不在其中：整节被清洗掉了（见 test_markdown_version_section_is_dropped）
    assert headings == [
        (2, "净销售额指标口径说明"),
        (2, "一、指标定义"),
        (2, "二、计算公式"),
    ]


def test_markdown_table_separator_is_not_data(tmp_path: Path) -> None:
    """`|---|---|` 是表头分隔行，**不能进数据行**。

    当成数据的话表头会变成 `---`，11.3 的「表头 + 当前行」直接错位——
    而错位的表读起来仍然"像"一张对的表，没人会去数。
    """
    parsed = parse_document(_write(tmp_path / "m.md", _MARKDOWN))

    table = parsed.blocks[1].table
    assert table is not None
    assert table.header == ("项目", "内容")
    assert table.rows == (("指标编码", "net_sales"), ("版本", "v1.2"))


def test_markdown_bold_line_before_table_becomes_caption(tmp_path: Path) -> None:
    """`**指标信息**` 紧接表格时是表名，不是正文。

    不认出来会多出一块 `… > 指标信息` 的 13 字垃圾块：不含信息、却参与召回、
    挤占 Top-K。而且表名挂上之后，**每一行都带着它**，行脱离表格也说得清自己是什么。
    """
    parsed = parse_document(_write(tmp_path / "m.md", _MARKDOWN))

    assert parsed.blocks[1].table is not None
    assert parsed.blocks[1].table.caption == "指标信息"
    assert all(b.text != "指标信息" for b in parsed.blocks)


def test_markdown_emphasis_is_stripped_but_text_kept(tmp_path: Path) -> None:
    parsed = parse_document(_write(tmp_path / "m.md", _MARKDOWN))

    body = " ".join(b.text for b in parsed.blocks if b.kind == "PARAGRAPH")
    assert "**" not in body
    assert "扣除让利与退回后" in body


def test_markdown_rule_lines_are_not_content(tmp_path: Path) -> None:
    """`---` 是分隔线。它会切出只剩横线的块，既进向量又毫无信息。"""
    parsed = parse_document(_write(tmp_path / "m.md", _MARKDOWN))

    assert all(set(b.text) != {"-"} for b in parsed.blocks)


def test_markdown_version_section_is_dropped(tmp_path: Path) -> None:
    """`六、版本变更记录` 整节丢弃（16.11.2 第 8 项），且**记进 dropped**。"""
    parsed = parse_document(_write(tmp_path / "m.md", _MARKDOWN))

    assert all("版本变更记录" not in b.text for b in parsed.blocks)
    assert any("版本变更记录" in item for item in parsed.dropped)


def test_markdown_has_no_page_numbers(tmp_path: Path) -> None:
    """Markdown 给不出页码，就给 None。**不用行号冒充**。

    行号会被下游当成页码过滤条件（"只要第 3 页"），静默筛掉一批 chunk。
    """
    parsed = parse_document(_write(tmp_path / "m.md", _MARKDOWN))

    assert parsed.page_count == 0
    assert all(b.page_no is None for b in parsed.blocks)


# ------------------------------------------------------------------ TXT


_TXT = """\
目标净销售额与达成率口径说明
----------------------------

【指标信息】
项目    内容
----  -------------------
指标编码  sales_target_amount
版本    v1.0

一、指标定义
------------

目标净销售额指公司按月度、区域、产品线下达的考核目标。
目标值由经营管理部会同财务部编制。

五、版本变更记录
----------------

【版本历史】
版本   日期  变更说明
v1.0  —   首次发布
"""


def test_txt_setext_heading_is_recognized(tmp_path: Path) -> None:
    parsed = parse_document(_write(tmp_path / "t.txt", _TXT))

    headings = [b.text for b in parsed.blocks if b.kind == "HEADING"]

    # `五、版本变更记录` 不在其中：整节被清洗掉（同 Markdown）
    assert headings == ["目标净销售额与达成率口径说明", "一、指标定义"]


def test_txt_table_separator_is_not_a_heading(tmp_path: Path) -> None:
    """TXT 里表格的对齐横线是 `----  ----------`（多段），**不是 setext 下划线**。

    区分靠"有没有内部空格"。不区分的话表格的表头行 `项目    内容`
    会被当成标题，整篇的标题路径跟着错位——而路径错了不会报错，
    只会让每个 chunk 都挂着一个假标题。
    """
    parsed = parse_document(_write(tmp_path / "t.txt", _TXT))

    headings = [b.text for b in parsed.blocks if b.kind == "HEADING"]
    assert "项目    内容" not in headings


def test_txt_table_is_downgraded_to_paragraphs(tmp_path: Path) -> None:
    """TXT 的表格**不解析结构**（11.2 只要求按段落切分）。

    固定宽度表格的列边界要靠字符对齐去猜，猜错的表现是单元格错位——
    比"降级成段落"糟得多。这条用例把这个取舍钉住，
    免得日后有人以为 TXT 也有表格结构。
    """
    parsed = parse_document(_write(tmp_path / "t.txt", _TXT))

    assert all(b.kind != "TABLE" for b in parsed.blocks)


def test_txt_paragraphs_are_joined_without_spaces(tmp_path: Path) -> None:
    """中文段落拼行**不插空格**。

    插了会切出 `sales_target_amount 版本` 这类原文里不存在的 token 形态。
    """
    parsed = parse_document(_write(tmp_path / "t.txt", _TXT))

    body = " ".join(b.text for b in parsed.blocks if b.kind == "PARAGRAPH")
    assert "考核目标。目标值由经营管理部" in body


# ------------------------------------------------------------------ DOCX


def _build_docx(path: Path) -> Path:
    """造一份结构真实的 DOCX：页眉页脚 + 多级标题 + 表格 + 表名前一行。"""
    document = docx.Document()
    document.styles["Normal"].font.size = Pt(10.5)
    section = document.sections[0]
    section.header.paragraphs[0].text = "示例科技股份有限公司　退货金额口径说明　[INTERNAL]"
    section.footer.paragraphs[0].text = "文件编号：MD-003　版本：v1.1"

    document.add_heading("退货金额口径说明", level=1)
    document.add_paragraph("表：指标信息")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"
    row = table.add_row().cells
    row[0].text = "指标编码"
    row[1].text = "return_amount"
    document.add_heading("一、指标定义", level=1)
    document.add_paragraph("退货金额指当期确认的退货与质量索赔冲减合计。")
    document.save(str(path))
    return path


def test_docx_heading_levels_come_from_styles(tmp_path: Path) -> None:
    parsed = parse_document(_build_docx(tmp_path / "d.docx"))

    headings = [(b.level, b.text) for b in parsed.blocks if b.kind == "HEADING"]
    assert headings == [(1, "退货金额口径说明"), (1, "一、指标定义")]


def test_docx_header_footer_are_dropped(tmp_path: Path) -> None:
    """DOCX 的页眉页脚在文档结构里是独立部件，**不用猜**（对比 PDF）。"""
    parsed = parse_document(_build_docx(tmp_path / "d.docx"))

    blob = parsed.text
    assert "示例科技股份有限公司" not in blob
    assert "文件编号" not in blob
    assert any("示例科技股份有限公司" in item for item in parsed.dropped)


def test_docx_table_keeps_caption_and_header(tmp_path: Path) -> None:
    parsed = parse_document(_build_docx(tmp_path / "d.docx"))

    table = next(b.table for b in parsed.blocks if b.kind == "TABLE")
    assert table is not None
    assert table.caption == "指标信息"
    assert table.header == ("项目", "内容")
    assert table.rows == (("指标编码", "return_amount"),)


def test_docx_body_order_is_preserved(tmp_path: Path) -> None:
    """段落与表格必须按**正文顺序**出现。

    python-docx 的 `document.paragraphs` 与 `document.tables` 是两个独立的列表，
    照着它们分别读会把所有表格堆到正文末尾——表格就落到错误的标题路径下了，
    而「这个数字属于哪一节」正是路径要回答的。
    """
    parsed = parse_document(_build_docx(tmp_path / "d.docx"))

    kinds = [b.kind for b in parsed.blocks]
    assert kinds == ["HEADING", "TABLE", "HEADING", "PARAGRAPH"]


def test_docx_has_no_page_numbers(tmp_path: Path) -> None:
    """DOCX 没有稳定页概念（分页由渲染器决定），给 None 而不是编一个。"""
    parsed = parse_document(_build_docx(tmp_path / "d.docx"))

    assert parsed.page_count == 0
    assert all(b.page_no is None for b in parsed.blocks)


# ------------------------------------------------------------------ PDF


def _build_pdf(path: Path, *, body_rows: int, pages: int = 1) -> Path:
    """造一份带页眉页脚、且表格会跨页的 PDF。

    与 `scripts/gen_corpus.py` 的排版路径一致：页眉页脚用 `on_page` 画、
    表格用 `repeatRows=1` 让它在续页重复表头——
    16.11.2 的「跨页表格」缺陷就是这么来的，解析层能不能还原表头，
    必须对着**真的跨了页**的 PDF 验。
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        PageTemplate,
        Paragraph,
        Spacer,
        TableStyle,
    )
    from reportlab.platypus import (
        Table as RLTable,
    )

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    font = "STSong-Light"
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=17)
    caption = ParagraphStyle("cap", fontName=font, fontSize=9.5)

    def on_page(canv: Any, doc: Any) -> None:
        canv.saveState()
        canv.setFont(font, 8)
        # 页眉：公司与标题在**同一行**（左、右各画一段），与真实语料一致
        canv.drawString(20 * mm, A4[1] - 14 * mm, "示例科技股份有限公司")
        canv.drawRightString(A4[0] - 20 * mm, A4[1] - 14 * mm, "退货金额口径说明　[INTERNAL]")
        # 页脚：文件编号（固定）与页码（每页不同）
        canv.drawString(20 * mm, 12 * mm, "文件编号：MD-003　版本：v1.1")
        canv.drawRightString(A4[0] - 20 * mm, 12 * mm, f"第 {doc.page} 页")
        canv.restoreState()

    document = BaseDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=22 * mm,
        bottomMargin=22 * mm,
    )
    document.addPageTemplates(
        [
            PageTemplate(
                id="main",
                frames=[
                    Frame(
                        document.leftMargin, document.bottomMargin, document.width, document.height
                    )
                ],
                onPage=on_page,
            )
        ]
    )

    story: list[Any] = [
        Paragraph(
            "退货金额口径说明", ParagraphStyle("title", fontName=font, fontSize=17, leading=24)
        )
    ]
    for page in range(pages):
        story.append(Paragraph("一、指标定义", body))
        story.append(Paragraph("退货金额指当期确认的退货与质量索赔冲减合计。", body))
        story.append(Paragraph("表：分区退货明细", caption))
        data = [["区域", "退货金额（万元）"]]
        data.extend([f"区域{page}-{i}", f"{i}.00"] for i in range(body_rows))
        table = RLTable(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    # **必须显式指定 CJK 字体**：reportlab 表格默认用 Helvetica，
                    # 它没有中文字形，会把每个汉字渲染成 `n`——
                    # 于是解析出来的是 `('nn', 'nnnnnnnn')`，而这条用例
                    # 想验的"表头被正确还原"就变成了验一串 n。
                    ("FONTNAME", (0, 0), (-1, -1), font),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 10))
    document.build(story)
    return path


def test_pdf_header_and_footer_are_dropped(tmp_path: Path) -> None:
    """页眉靠"跨页重复"、页脚靠"跨页重复 + 页码特征"认出。

    页脚每页的页码都不同：不先把数字归一成 `#` 再比，**一页都匹配不上**，
    而症状只是"页脚进了 chunk"，不报错、不告警。
    """
    parsed = parse_document(_build_pdf(tmp_path / "p.pdf", body_rows=10, pages=2))

    blob = parsed.text
    assert "示例科技股份有限公司" not in blob
    assert "第 1 页" not in blob and "第 2 页" not in blob
    assert any("文件编号" in item for item in parsed.dropped)


def test_pdf_keeps_page_numbers(tmp_path: Path) -> None:
    parsed = parse_document(_build_pdf(tmp_path / "p.pdf", body_rows=50))

    # 页数不写死：写死会让这条用例在版面参数微调后变成"测试坏了"而不是"页码坏了"
    assert parsed.page_count >= 2
    assert {b.page_no for b in parsed.blocks} == set(range(1, parsed.page_count + 1))


def test_pdf_caption_is_attached_to_the_table(tmp_path: Path) -> None:
    """`表：分区退货明细` 在 PDF 里是**独立一行文本**，表格是另一个对象。

    不跨对象把它挂过去，每张表就都没有表名，
    11.3 的「每块保留表名」落空，而表名正是"这块数据是什么"的答案。
    """
    parsed = parse_document(_build_pdf(tmp_path / "p.pdf", body_rows=10))

    table = next(b.table for b in parsed.blocks if b.kind == "TABLE")
    assert table is not None
    assert table.caption == "分区退货明细"
    assert table.header == ("区域", "退货金额（万元）")


def test_pdf_cross_page_table_keeps_header_and_caption(tmp_path: Path) -> None:
    """**跨页表格的表头还原**（16.11.2 第 7 项）。

    表格撑破一页时续页会重复表头，`pdfplumber` 因此给出两张"表头相同、
    表名缺失"的独立表格。续接表必须继承表名，否则第 2 页那些行只剩一行裸数据，
    谁也说不清它属于哪张表。

    两条断言缺一不可：**表头在**（证明跨页没有把表头丢掉）、
    **表名也在**（证明续接行能被追溯到它所属的表）。
    """
    parsed = parse_document(_build_pdf(tmp_path / "p.pdf", body_rows=45, pages=2))

    tables = [b.table for b in parsed.blocks if b.kind == "TABLE" and b.table is not None]
    assert len(tables) >= 3, "这份 PDF 应当至少有跨页拆出的 3 张表"

    headers = {" | ".join(t.header) for t in tables}
    assert len(headers) == 1, f"续页表头应当一致，实际 {headers}"
    assert all(t.caption == "分区退货明细" for t in tables), [t.caption for t in tables]
    # 每张表都真的有数据行——只有表头的空表说明分带把内容切丢了
    assert all(t.rows for t in tables)


def test_scanned_pdf_reports_missing_text_layer(tmp_path: Path) -> None:
    """图片型 PDF **不抛异常**：解析层如实报告，由入库层记 FAILED（11.2）。

    抛异常会把"这份文件我们读不了"变成"解析器坏了"，
    而 16.11.2 指定的期望行为是 `status = FAILED` + `error_summary`。
    """
    parsed = parse_document(_build_scanned_pdf(tmp_path / "s.pdf"))

    assert parsed.has_text_layer is False
    assert parsed.blocks == ()
    assert parsed.page_count >= 1


def _build_scanned_pdf(path: Path) -> Path:
    """图片型 PDF：只有位图，没有文本层。"""
    import io

    from PIL import Image
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as rl_canvas

    buffer = io.BytesIO()
    Image.new("RGB", (800, 1000), "white").save(buffer, format="JPEG")
    buffer.seek(0)
    out = rl_canvas.Canvas(str(path), pagesize=A4)
    out.drawImage(ImageReader(buffer), 0, 0, width=A4[0], height=A4[1])
    out.showPage()
    out.save()
    return path


# ------------------------------------------------------------------ 清洗规则（纯函数）


def test_repeated_ignores_page_numbers() -> None:
    """页码每页不同，必须归一化后再比。"""
    repeated = _repeated(["文件编号：A 第 1 页", "文件编号：A 第 2 页"])

    assert repeated == {"文件编号：A 第 # 页"}


def test_repeated_needs_two_occurrences() -> None:
    assert _repeated(["只出现一次的行"]) == set()


def test_short_document_header_is_recognized_by_classification_mark() -> None:
    """**单页文档**没有"跨页重复"可依，密级标记是它唯一的结构特征。

    不认它，语料里那 3 份单页 PDF 的页眉就会原样进 chunk。
    """
    matched, samples = _detect_page_furniture(["示例科技股份有限公司 A 文件 [INTERNAL]\n正文"])

    assert matched == {"示例科技股份有限公司 A 文件 [INTERNAL]"}
    assert samples == ["示例科技股份有限公司 A 文件 [INTERNAL]"]


def test_single_page_footer_is_recognized_by_page_number() -> None:
    matched, _ = _detect_page_furniture(["正文\n文件编号：A 版本：v1.0 第 1 页"])

    assert matched == {"文件编号：A 版本：v#.# 第 # 页"}


def test_first_line_without_header_features_is_kept() -> None:
    """首行没有页眉特征时**不能当页眉丢**。

    单页文档的第一行往往就是正文开头，见行就丢等于吃掉正文。
    """
    matched, _ = _detect_page_furniture(["净销售额指本期实现的收入。\n正文继续"])

    assert matched == set()


def test_inherit_captions_requires_adjacent_identical_header() -> None:
    """续接表继承表名，前提是**紧邻且表头相同**。

    中间隔着正文段落时，下一张表可能是另一张表——从宽继承会挂错表名，
    而错表名比没有表名更难发现。
    """
    header = ("区域", "金额")
    first = Block(
        kind="TABLE", table=TableBlock(caption="分区明细", header=header, rows=(("华东", "1"),))
    )
    continuation = Block(
        kind="TABLE", table=TableBlock(caption=None, header=header, rows=(("华南", "2"),))
    )
    other_header = Block(kind="TABLE", table=TableBlock(caption=None, header=("渠道", "金额")))
    paragraph = Block(kind="PARAGRAPH", text="中间还有一段正文")

    inherited = _inherit_captions([first, continuation, other_header])
    assert inherited[1].table is not None and inherited[1].table.caption == "分区明细"
    assert inherited[2].table is not None and inherited[2].table.caption is None

    separated = _inherit_captions([first, paragraph, continuation])
    assert separated[2].table is not None and separated[2].table.caption is None


def test_parsed_document_text_excludes_tables() -> None:
    """`ParsedDocument.text` 只拼正文：它是给人看的纯文本，
    表格的序列化形态由分块层决定（见 `chunker.serialize_table`）。"""
    parsed = ParsedDocument(
        blocks=(
            Block(kind="HEADING", text="标题", level=1),
            Block(kind="TABLE", table=TableBlock(header=("a",), rows=(("b",),))),
            Block(kind="PARAGRAPH", text="正文"),
        )
    )

    assert parsed.text == "标题\n正文"
